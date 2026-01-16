import streamlit as st
import pandas as pd
import re
from rapidfuzz import fuzz
from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from io import BytesIO
import random

# ============ دوال المساعدة (Helper Functions) ============

def generate_color():
    """توليد ألوان للمجموعات المتشابهة"""
    colors = ["FFB3BA", "BAFFC9", "BAE1FF", "FFFFBA", "FFDFBA", "E0BBE4", "D4A5A5", "A8E6CF"]
    return random.choice(colors)

def normalize_name(name):
    """تطبيع النصوص العربية لضمان دقة المطابقة"""
    if pd.isnull(name): return ""
    name = str(name).strip()
    name = name.replace("ه", "ة").replace("أ", "ا").replace("إ", "ا").replace("آ", "ا")
    name = name.replace("ى", "ي").replace("ئ", "ي")
    name = re.sub(r'(عبد)([^\s])', r'\1 \2', name)
    return " ".join(name.split()).lower()

def get_first_three_words(name):
    """استخراج أول 3 كلمات من الاسم"""
    words = str(name).split()
    return " ".join(words[:3]) if len(words) >= 3 else " ".join(words)

def is_first_three_words_match(name1, name2):
    """التأكد من أن أول 3 كلمات متطابقة تماماً"""
    words1 = name1.split()
    words2 = name2.split()
    length = min(len(words1), len(words2), 3)
    if length == 0: return False
    return all(words1[i] == words2[i] for i in range(length))

# ============ منطق المطابقة (Matching Logic) ============

def find_best_match_in_db(normalized_name, database_map, threshold=85):
    """البحث عن أفضل مطابقة في خريطة بيانات معينة"""
    best_match = None
    best_score = 0
    
    for db_name in database_map.keys():
        score = fuzz.ratio(normalized_name, db_name)
        if score > best_score:
            best_score = score
            best_match = db_name

    if best_score >= threshold and best_match:
        if is_first_three_words_match(normalized_name, best_match) or best_match.startswith(normalized_name):
            return best_match, best_score
            
    for db_name in database_map.keys():
        if db_name.startswith(normalized_name) or normalized_name.startswith(db_name):
            return db_name, fuzz.ratio(normalized_name, db_name)
            
    return None, 0

def match_names_core(names_df, db_df, col_file, col_db, selected_cols):
    """الدالة الأساسية لمطابقة ملف مع قاعدة بيانات واحدة"""
    db_df = db_df.copy()
    db_df["norm"] = db_df[col_db].apply(normalize_name)
    db_map = db_df.drop_duplicates(subset=["norm"]).set_index("norm").to_dict(orient="index")
    
    results = []
    for _, row in names_df.iterrows():
        orig = row[col_file]
        norm = normalize_name(orig)
        best_m, score = find_best_match_in_db(norm, db_map)
        
        res = {"الاسم الأصلي": orig, "الاسم المطابق": "", "نسبة التطابق": "", "ملاحظة": "❌ لا يوجد تطابق"}
        if best_m:
            data = db_map[best_m]
            res.update({"الاسم المطابق": data[col_db], "نسبة التطابق": f"{round(score)}%", "ملاحظة": "✅ تطابق"})
            for c in selected_cols: res[c] = data.get(c, "")
        else:
            for c in selected_cols: res[c] = ""
        results.append(res)
    return pd.DataFrame(results)

# ============ التصدير إلى إكسل ============

def to_excel_styled(df, highlight_col="ملاحظة"):
    output = BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df.to_excel(writer, index=False)
        ws = writer.book.active
        red_fill = PatternFill(start_color="FFCCCC", end_color="FFCCCC", fill_type="solid")
        
        for column in ws.columns:
            max_length = max(len(str(cell.value)) if cell.value else 0 for cell in column)
            ws.column_dimensions[column[0].column_letter].width = max_length + 2
        
        col_idx = None
        for i, cell in enumerate(ws[1], 1):
            if cell.value == highlight_col: col_idx = i; break
            
        if col_idx:
            for row in ws.iter_rows(min_row=2):
                if row[col_idx-1].value and "❌" in str(row[col_idx-1].value):
                    for cell in row: cell.fill = red_fill
    output.seek(0)
    return output

# ============ واجهة التطبيق (UI) ============

st.set_page_config(page_title="نظام المطابقة الموحد", layout="wide")
st.markdown("<h3 style='text-align: right;'>👨‍💻 برمجة: محمد عبدالجليل</h3>", unsafe_allow_html=True)
st.title("🔐 نظام مطابقة الأسماء والرواتب المتقدم")

password = st.sidebar.text_input("أدخل كلمة المرور:", type="password")

if password == "mjaleel":
    tabs = st.tabs(["📋 مطابقة عامة", "🏢 مطابقة أقسام", "🔄 مطابقة ثلاثية", "🏦 مطابقة الرواتب (قاعدتين/ورقتين)"])

    # --- التبويب 1 و 2 (تم دمج المنطق للاختصار) ---
    for i, tab in enumerate([tabs[0], tabs[1]]):
        with tab:
            st.subheader("📋 مطابقة ملف مع قاعدة بيانات" if i==0 else "🏢 مطابقة الأقسام")
            c1, c2 = st.columns(2)
            f1 = c1.file_uploader(f"رفع ملف الأسماء ({i})", type="xlsx")
            f2 = c2.file_uploader(f"رفع قاعدة البيانات ({i})", type="xlsx")
            
            if f1 and f2:
                df1, df2 = pd.read_excel(f1), pd.read_excel(f2)
                col_f = st.selectbox("عمود الاسم في ملفك:", df1.columns, key=f"f{i}")
                col_d = st.selectbox("عمود الاسم في القاعدة:", df2.columns, key=f"d{i}")
                extra = st.multiselect("أعمدة إضافية لجلبها:", [c for c in df2.columns if c != col_d], key=f"e{i}")
                
                if st.button("🚀 بدء المطابقة", key=f"b{i}"):
                    res = match_names_core(df1, df2, col_f, col_d, extra)
                    st.dataframe(res, use_container_width=True)
                    st.download_button("⬇️ تحميل", to_excel_styled(res), "results.xlsx")

    # --- التبويب 4: البحث المتسلسل (المطلوب برمجياً) ---
    with tabs[3]:
        st.subheader("🏦 مطابقة الرواتب والـ IBAN (بحث متسلسل)")
        st.info("المنطق: يبحث في القاعدة 1 أولاً، وإذا لم يجد ينتقل للقاعدة 2 (الرافدين).")
        
        mode = st.radio("مصدر القواعد:", ["ورقتين في ملف واحد", "ملفين منفصلين"], horizontal=True)
        
        col1, col2, col3 = st.columns(3)
        file_main = col1.file_uploader("📄 ملف الأسماء الأصلي", type="xlsx", key="main_pay")
        
        db1_df, db2_df = None, None
        
        if mode == "ورقتين في ملف واحد":
            combined_file = col2.file_uploader("📊 ملف القواعد الموحد", type="xlsx")
            if combined_file:
                xl = pd.ExcelFile(combined_file)
                sh1 = col2.selectbox("ورقة (الموطنة):", xl.sheet_names)
                sh2 = col3.selectbox("ورقة (الرافدين):", xl.sheet_names, index=min(1, len(xl.sheet_names)-1))
                db1_df = pd.read_excel(combined_file, sheet_name=sh1)
                db2_df = pd.read_excel(combined_file, sheet_name=sh2)
        else:
            f_db1 = col2.file_uploader("📊 ملف قاعدة 1", type="xlsx")
            f_db2 = col3.file_uploader("🏛️ ملف قاعدة 2 (الرافدين)", type="xlsx")
            if f_db1 and f_db2:
                db1_df, db2_df = pd.read_excel(f_db1), pd.read_excel(f_db2)

        if file_main and db1_df is not None and db2_df is not None:
            df_main = pd.read_excel(file_main)
            st.markdown("---")
            ec1, ec2, ec3 = st.columns(3)
            c_main = ec1.selectbox("عمود الاسم (الأصلي):", df_main.columns)
            c_db1 = ec2.selectbox("عمود الاسم (قاعدة 1):", db1_df.columns)
            c_db2 = ec3.selectbox("عمود الاسم (قاعدة 2):", db2_df.columns)
            
            sel1 = st.multiselect("بيانات من قاعدة 1 (IBAN):", [c for c in db1_df.columns if c != c_db1])
            sel2 = st.multiselect("بيانات من قاعدة 2 (IBAN):", [c for c in db2_df.columns if c != c_db2])
            
            if st.button("🚀 تشغيل مطابقة الرواتب", type="primary"):
                # تجهيز القواعد
                db1_df["norm"] = db1_df[c_db1].apply(normalize_name)
                db2_df["norm"] = db2_df[c_db2].apply(normalize_name)
                map1 = db1_df.drop_duplicates(subset=["norm"]).set_index("norm").to_dict(orient="index")
                map2 = db2_df.drop_duplicates(subset=["norm"]).set_index("norm").to_dict(orient="index")
                
                final_results = []
                for _, row in df_main.iterrows():
                    name = row[c_main]
                    norm = normalize_name(name)
                    
                    entry = {"الاسم الأصلي": name, "الاسم المطابق": "", "المصدر": "❌ لم يتم العثور", "ملاحظة": "❌"}
                    for c in sel1: entry[f"[ق1] {c}"] = ""
                    for c in sel2: entry[f"[ق2] {c}"] = ""
                    
                    # محاولة قاعدة 1
                    bm1, s1 = find_best_match_in_db(norm, map1)
                    if bm1:
                        d = map1[bm1]
                        entry.update({"الاسم المطابق": d[c_db1], "المصدر": "✅ القاعدة 1", "ملاحظة": "✅"})
                        for c in sel1: entry[f"[ق1] {c}"] = d.get(c, "")
                    else:
                        # محاولة قاعدة 2
                        bm2, s2 = find_best_match_in_db(norm, map2)
                        if bm2:
                            d = map2[bm2]
                            entry.update({"الاسم المطابق": d[c_db2], "المصدر": "✅ قاعدة الرافدين", "ملاحظة": "✅"})
                            for c in sel2: entry[f"[ق2] {c}"] = d.get(c, "")
                    
                    final_results.append(entry)
                
                res_df = pd.DataFrame(final_results)
                st.dataframe(res_df, use_container_width=True)
                st.download_button("⬇️ تحميل كشف الرواتب", to_excel_styled(res_df), "Salary_Match.xlsx")

elif password:
    st.error("❌ كلمة المرور خاطئة")
else:
    st.warning("الرجاء إدخال كلمة المرور")
 
